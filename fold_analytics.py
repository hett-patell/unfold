"""Local enrichment + analytics for Fold data (no network; reads fold.db)."""
import datetime as dt
import json
import re
import sqlite3
from collections import defaultdict
from pathlib import Path

# VPA handle -> payment app (display flavour only)
UPI_APP = {
    "ybl": "PhonePe", "ibl": "PhonePe", "axl": "PhonePe", "apl": "Amazon Pay",
    "yapl": "Amazon Pay", "paytm": "Paytm", "ptys": "Paytm", "ptm": "Paytm",
    "oksbi": "GPay", "okhdfcbank": "GPay", "okaxis": "GPay", "okicici": "GPay",
    "fam": "FamPay", "jupiteraxis": "Jupiter", "superyes": "Super.money",
}

# ordered keyword -> category rules (first match wins)
_RULES = [
    (r"swiggy|zomato|eatsure|faasos|dominos|pizza|kfc|mcdonald|burger\s?king|"
     r"restaurant|dhaba|cafe|bakery|baskin|starbucks|chai|tea\b|biryani|kitchen|foods?\b", "Food & Drinks"),
    (r"instamart|blinkit|zepto|bigbasket|grofers|dmart|jiomart|dunzo|grocer|kirana|supermarket|mart\b", "Groceries"),
    (r"\buber\b|\bola\b|rapido|irctc|redbus|makemytrip|ixigo|goibibo|indigo|vistara|"
     r"petrol|fuel|hpcl|iocl|bpcl|fastag|parking|metro|namma|\btoll\b|railway|travels?\b", "Transport"),
    (r"myntra|ajio|nykaa|zudio|flipkart|meesho|tatacliq|lifestyle|decathlon|"
     r"reliance.?trends|\bzara\b|shoppers|amazon(?!.?pay)", "Shopping"),
    (r"netflix|spotify|hotstar|jiocinema|sonyliv|zee5|youtube|prime\s?video|"
     r"apple\.com/bill|icloud|google.*(storage|one)|chatgpt|openai|claude|notion|adobe|canva|subscription", "Subscription"),
    (r"\bjio\b|airtel|vodafone|\bvi\b|bsnl|recharge|\bdth\b|tatasky|broadband|fibernet|wifi", "Bill"),
    (r"electric|torrent.?power|adani|mseb|bescom|tneb|indane|gas\b|water.?bill|municipal|\brent\b", "Bill"),
    (r"salary|payroll|stipend", "Salary"),
    (r"\batm\b|cash.?w(d|ithdraw)|nwd/|cwdr", "Cash Withdrawal"),
    (r"groww|zerodha|kuvera|indmoney|upstox|smallcase|coin\b|mfautopay|\bsip\b|mutual|"
     r"\bnps\b|\bppf\b|\bfd\b|recurring.?dep", "Investment"),
    (r"\bcred\b|creditcard|cc.?pay|billdesk|card.?payment", "Credit Bill"),
    (r"insurance|policy|\blic\b|premium|hdfc.?life|max.?life", "Insurance"),
    (r"pharm|1mg|apollo|medplus|hospital|clinic|diagnostic|\blab\b|netmeds|medic", "Medical"),
    (r"amazon.?pay|\bwallet\b top", "Top-up"),
]
_RULES = [(re.compile(p, re.I), c) for p, c in _RULES]
NON_SPEND = ("Self Transfer", "Cash Withdrawal")


def extract_counterparty(narration, merchant_name=None):
    if merchant_name:
        return merchant_name
    n = (narration or "").strip()
    if not n:
        return None
    up = n.upper()
    if up.startswith("UPI/"):
        parts = n.split("/")
        if len(parts) >= 5 and parts[4].strip():
            return parts[4].strip()
    if re.match(r"(NEFT|RTGS|IMPS)", up):                 # "NEFT CR-IFSC-NAME-..."
        segs = [s.strip() for s in n.split("-") if s.strip()]
        if len(segs) >= 3:
            return segs[2]
    if up.startswith("UPI-"):                             # "UPI-NAME-VPA-..."
        segs = [s.strip() for s in n.split("-") if s.strip()]
        if len(segs) >= 2:
            return segs[1]
    return None


def app_for(counterparty):
    if not counterparty:
        return None
    m = re.search(r"@([a-z]+)$", counterparty.lower()) or re.search(r"([a-z]+)$", counterparty.lower())
    return UPI_APP.get(m.group(1)) if m else None


def guess_category(narration, counterparty, ttype):
    text = f"{narration or ''} {counterparty or ''}".lower()
    for rx, cat in _RULES:
        if rx.search(text):
            return cat
    return "Income (other)" if ttype == "INCOMING" else "Uncategorized"


def apply_aliases(counterparty, narration, aliases):
    hay = f"{counterparty or ''} {narration or ''}".lower()
    for pat, name, cat in aliases or ():
        if pat and pat in hay:
            return (name or None), (cat or None)
    return None, None


def load_aliases(path):
    pp = Path(path) if path else None
    if not pp or not pp.exists():
        return []
    try:
        data = json.loads(pp.read_text())
    except Exception:
        return []
    return [(k.lower(), v.get("name"), v.get("category"))
            for k, v in data.items() if not k.startswith("_") and isinstance(v, dict)]


def enrich(narration, merchant_name, category_name, ttype, aliases=None):
    cp = extract_counterparty(narration, merchant_name)
    alias_name, alias_cat = apply_aliases(cp, narration, aliases)
    if alias_name:
        cp = alias_name
    cat = alias_cat or category_name or guess_category(narration, cp, ttype)
    return cp, cat


# ---- analytics -------------------------------------------------------------
def _rows(db, aliases=None):
    cur = db.execute("SELECT txn_date, type, mode, amount, category_name, "
                     "merchant_name, narration, excluded_from_cash_flow FROM transactions")
    out = []
    for date, typ, mode, amt, cat, merch, narr, excl in cur:
        cp, category = enrich(narr, merch, cat, typ, aliases)
        out.append({"date": date or "", "month": (date or "")[:7], "type": typ,
                    "amount": amt or 0, "mode": mode or "OTHER", "category": category,
                    "counterparty": cp, "narration": narr, "excluded": excl, "fold_cat": cat})
    return out


def _load(db_path, aliases_path=None):
    aliases = load_aliases(aliases_path)
    db = sqlite3.connect(str(db_path))
    rows = _rows(db, aliases)
    accounts = [json.loads(r[0]) for r in db.execute("SELECT raw FROM accounts")]
    db.close()
    return rows, accounts


def _spendable(r):
    return not r["excluded"] and r["category"] not in NON_SPEND


def compute(db_path, month=None, aliases_path=None):
    rows, accounts = _load(db_path, aliases_path)

    months = sorted({r["month"] for r in rows if r["month"]})
    month = month or (months[-1] if months else dt.date.today().strftime("%Y-%m"))

    # coverage: how much did enrichment help
    n = len(rows)
    fold_unc = sum(1 for r in rows if not r["fold_cat"])
    enr_unc = sum(1 for r in rows if r["category"] == "Uncategorized")

    # monthly income/expense
    monthly = {}
    for r in rows:
        if not r["month"] or not _spendable(r):
            continue
        m = monthly.setdefault(r["month"], {"income": 0.0, "expense": 0.0})
        if r["type"] == "INCOMING":
            m["income"] += r["amount"]
        elif r["type"] == "OUTGOING":
            m["expense"] += r["amount"]
    monthly_series = [{"month": m, **{k: round(v, 2) for k, v in monthly[m].items()},
                       "net": round(monthly[m]["income"] - monthly[m]["expense"], 2)}
                      for m in months]

    def agg(field, month_filter=None, typ="OUTGOING"):
        d = defaultdict(lambda: [0.0, 0])
        for r in rows:
            if r["type"] != typ or not _spendable(r):
                continue
            if month_filter and r["month"] != month_filter:
                continue
            key = r[field] or ("Uncategorized" if field == "category" else "(unknown)")
            d[key][0] += r["amount"]
            d[key][1] += 1
        return sorted(([k, round(v[0], 2), v[1]] for k, v in d.items()),
                      key=lambda x: -x[1])

    # recurring: same counterparty seen in >=3 distinct months
    by_cp_month = defaultdict(set)
    by_cp_stats = defaultdict(lambda: [0.0, 0])
    for r in rows:
        if r["type"] != "OUTGOING" or not r["counterparty"] or not _spendable(r):
            continue
        by_cp_month[r["counterparty"]].add(r["month"])
        by_cp_stats[r["counterparty"]][0] += r["amount"]
        by_cp_stats[r["counterparty"]][1] += 1
    recurring = sorted(
        ([cp, len(m), round(by_cp_stats[cp][0], 2), by_cp_stats[cp][1]]
         for cp, m in by_cp_month.items() if len(m) >= 3),
        key=lambda x: (-x[1], -x[2]))

    biggest = sorted(
        ([r["date"], r["counterparty"] or r["narration"][:30], r["category"], round(r["amount"], 2)]
         for r in rows if r["type"] == "OUTGOING" and _spendable(r)),
        key=lambda x: -x[3])[:10]

    total_in = round(sum(r["amount"] for r in rows if r["type"] == "INCOMING" and _spendable(r)), 2)
    total_out = round(sum(r["amount"] for r in rows if r["type"] == "OUTGOING" and _spendable(r)), 2)

    return {
        "month": month, "months": months, "n": n,
        "date_range": [rows and min(r["date"] for r in rows if r["date"]),
                       rows and max(r["date"] for r in rows if r["date"])],
        "coverage": {"total": n, "fold_uncat": fold_unc, "enriched_uncat": enr_unc,
                     "recovered": fold_unc - enr_unc},
        "balances": [{"bank": (a.get("financial_information_provider") or {}).get("name"),
                      "masked": a.get("masked_account_number"),
                      "balance": a.get("current_balance")} for a in accounts],
        "total_balance": round(sum((a.get("current_balance") or 0) for a in accounts), 2),
        "lifetime": {"income": total_in, "expense": total_out, "net": round(total_in - total_out, 2)},
        "monthly": monthly_series,
        "cat_month": agg("category", month),
        "merch_month": agg("counterparty", month),
        "cat_all": agg("category"),
        "merch_all": agg("counterparty"),
        "by_mode": spend_by_mode(rows),
        "forecast": forecast_month(rows),
        "recurring": recurring[:15],
        "biggest": biggest,
    }


def parse_investments(inv, bank_total):
    agg = inv.get("aggregation") or {}
    mf = inv.get("mutual_funds") or {}
    demat = agg.get("demat_holdings") or {}
    stocks_total = demat.get("total_value") or 0
    mf_total = (agg.get("mutual_fund_stats") or {}).get("total_current_value") or mf.get("total_current_value") or 0
    fd_total = (agg.get("fixed_deposits_stats") or {}).get("total_maturity_amount") or 0
    nps_total = (agg.get("nps_stats") or {}).get("total_value") or 0
    inv_total = round(stocks_total + mf_total + fd_total + nps_total, 2)
    ar = mf.get("absolute_returns") or {}
    schemes = [{"name": x.get("name"), "value": x.get("value"), "invested": x.get("invested"),
                "gain_loss": x.get("gain_loss"), "returns": x.get("returns"), "xirr": x.get("xirr"),
                "category": x.get("asset_sub_category") or x.get("asset_category")}
               for x in (mf.get("schemes") or [])]
    return {"generated": dt.datetime.now().isoformat(), "bank": round(bank_total, 2),
            "stocks": {"value": round(stocks_total, 2), "holdings": demat.get("total_holdings"),
                       "shares": demat.get("total_shares")},
            "mutual_funds": {"value": round(mf_total, 2), "xirr": mf.get("xirr"),
                             "benchmark_xirr": mf.get("benchmark_xirr"), "gain": ar.get("total_value"),
                             "gain_pct": ar.get("total_gain_percentage"), "schemes": schemes},
            "fd": round(fd_total, 2), "nps": round(nps_total, 2),
            "investments_total": inv_total, "net_worth": round(bank_total + inv_total, 2)}


def _investments_card(nw):
    mf, st = nw["mutual_funds"], nw["stocks"]
    rows = "".join(
        f'<tr><td>{_esc((s.get("name") or "")[:34])}</td><td class="n">{_esc(inr(s.get("value") or 0))}</td>'
        f'<td class="n">{("+" if (s.get("gain_loss") or 0)>=0 else "")}{_esc(inr(s.get("gain_loss") or 0))}</td>'
        f'<td class="n">{s.get("xirr") if s.get("xirr") is not None else "—"}</td></tr>'
        for s in mf["schemes"]) or '<tr><td colspan="4" class="empty">No schemes.</td></tr>'
    return (f'<div class="card"><h2>Investments <span class="m">· {_esc(inr_c(nw["investments_total"]))} total</span></h2>'
            '<div class="tiles" style="margin-bottom:14px">'
            f'<div class="tile"><div class="k">Stocks</div><div class="v">{_esc(inr_c(st["value"]))}</div></div>'
            f'<div class="tile"><div class="k">Mutual funds</div><div class="v">{_esc(inr_c(mf["value"]))}</div></div>'
            f'<div class="tile"><div class="k">MF XIRR</div><div class="v">{mf.get("xirr","—")}%</div></div>'
            f'<div class="tile"><div class="k">MF gain</div><div class="v">{_esc(inr_c(mf.get("gain") or 0))}</div></div></div>'
            '<table><thead><tr><th>Mutual fund</th><th class="n">Value</th><th class="n">Gain</th><th class="n">XIRR%</th></tr></thead>'
            f'<tbody>{rows}</tbody></table></div>')


# ---- report (local HTML dashboard) -----------------------------------------
import html as _html


def _ind_group(n):
    import re as _re
    n = int(round(n or 0)); sign = "-" if n < 0 else ""; s = str(abs(n))
    if len(s) <= 3:
        return sign + s
    head, tail = s[:-3], s[-3:]
    head = _re.sub(r"(?<=\d)(?=(\d\d)+$)", ",", head)
    return f"{sign}{head},{tail}"


def inr(n):
    return "₹" + _ind_group(n)


def hr(w=52, ch="─"):
    return ch * w


def table(headers, rows, aligns=None, indent="  ", gap="  ", header=True, underline=True):
    ncol = len(headers)
    widths = [len(str(headers[i])) if header else 0 for i in range(ncol)]
    for r in rows:
        for i in range(ncol):
            widths[i] = max(widths[i], len(str(r[i])))
    aligns = aligns or ["l"] * ncol

    def fmt(cells):
        parts = []
        for i in range(ncol):
            c = str(cells[i])
            parts.append(c.rjust(widths[i]) if aligns[i] == "r" else c.ljust(widths[i]))
        return (indent + gap.join(parts)).rstrip()

    lines = []
    if header:
        lines.append(fmt(headers))
        if underline:
            lines.append(indent + hr(sum(widths) + len(gap) * (ncol - 1)))
    lines += [fmt(r) for r in rows]
    return lines


def print_bars(items, width=22, indent="  ", label_w=18):
    if not items:
        print(indent + "(none)")
        return
    top = items[0][1] or 1
    nw = min(label_w, max(len(str(n)[:label_w]) for n, _a, *_ in items))
    aw = max(len(inr_c(a)) for _n, a, *_ in items)
    for name, amt, *rest in items:
        bar = "█" * max(1, round(width * (amt / top))) if amt else ""
        cnt = f" ({rest[0]})" if rest else ""
        print(f"{indent}{str(name)[:label_w]:<{nw}}  {inr_c(amt):>{aw}}  {bar}{cnt}")


def inr_c(n):
    n = n or 0
    a, s = abs(n), ("-" if n < 0 else "")
    if a >= 1e7:
        return f"{s}₹{a/1e7:.2f}Cr"
    if a >= 1e5:
        return f"{s}₹{a/1e5:.2f}L"
    if a >= 1e3:
        return f"{s}₹{a/1e3:.1f}k"
    return f"{s}₹{a:.0f}"


def _esc(x):
    return _html.escape(str(x))


def _monthly_svg(monthly):
    ms = monthly[-12:]
    if not ms:
        return '<p class="empty">No monthly data.</p>'
    W, H, padL, padR, padT, padB = 920, 300, 60, 12, 14, 46
    plotW, plotH = W - padL - padR, H - padT - padB
    mx = max([max(m["income"], m["expense"]) for m in ms]) or 1
    n = len(ms)
    gw = plotW / n
    bw = min(16, gw / 3.2)

    def yv(v):
        return padT + plotH - (v / mx) * plotH
    p = [f'<svg viewBox="0 0 {W} {H}" class="chart" role="img" aria-label="Monthly income vs expense">']
    for fr in (0, .25, .5, .75, 1):
        gy = padT + plotH - fr * plotH
        p.append(f'<line x1="{padL}" y1="{gy:.1f}" x2="{W-padR}" y2="{gy:.1f}" class="grid"/>')
        p.append(f'<text x="{padL-6}" y="{gy+3:.1f}" class="ytick">{_esc(inr_c(mx*fr))}</text>')
    for i, m in enumerate(ms):
        cx = padL + i * gw + gw / 2
        inc, exp = m["income"], m["expense"]
        xi, xe = cx - bw - 1, cx + 1
        p.append(f'<rect x="{xi:.1f}" y="{yv(inc):.1f}" width="{bw:.1f}" height="{padT+plotH-yv(inc):.1f}" rx="3" fill="var(--income)"><title>{m["month"]} · in {_esc(inr_c(inc))}</title></rect>')
        p.append(f'<rect x="{xe:.1f}" y="{yv(exp):.1f}" width="{bw:.1f}" height="{padT+plotH-yv(exp):.1f}" rx="3" fill="var(--expense)"><title>{m["month"]} · out {_esc(inr_c(exp))}</title></rect>')
        p.append(f'<text x="{cx:.1f}" y="{H-padB+16}" class="xtick">{_esc(m["month"][2:])}</text>')
        net = m["net"]
        p.append(f'<text x="{cx:.1f}" y="{H-padB+30}" class="xnet" fill="var({"--pos" if net>=0 else "--neg"})">{_esc(inr_c(net))}</text>')
    p.append("</svg>")
    return "".join(p)


def _hbars(items, color_var, max_rows=10):
    items = items[:max_rows]
    if not items:
        return '<p class="empty">No data this month.</p>'
    mx = max(v for _, v, *_ in items) or 1
    rowH, W, barX, valW = 30, 460, 150, 78
    barW = W - barX - valW
    H = len(items) * rowH + 6
    p = [f'<svg viewBox="0 0 {W} {H}" class="hb" role="img">']
    for i, (label, val, *rest) in enumerate(items):
        cy = i * rowH + rowH / 2 + 3
        w = max(2, (val / mx) * barW)
        lab = _esc(str(label))[:20]
        cnt = f" · {rest[0]} txns" if rest else ""
        p.append(f'<text x="0" y="{cy+4:.0f}" class="hbl">{lab}</text>')
        p.append(f'<rect x="{barX}" y="{cy-8:.0f}" width="{w:.1f}" height="16" rx="4" fill="var({color_var})"><title>{lab}: {_esc(inr(val))}{cnt}</title></rect>')
        p.append(f'<text x="{barX+w+6:.1f}" y="{cy+4:.0f}" class="hbv">{_esc(inr_c(val))}</text>')
    p.append("</svg>")
    return "".join(p)


_CSS = """
*{box-sizing:border-box} body{margin:0}
.wrap{max-width:1000px;margin:0 auto;padding:28px 20px 60px}
.viz{color-scheme:light;
  --bg:#f4f3f0;--surface-1:#fcfcfb;--border:#e7e6e1;--text-primary:#0b0b0b;
  --text-secondary:#52514e;--muted:#8a8983;--grid:#eceae5;
  --income:#2a78d6;--expense:#eb6834;--cat:#256abf;--cp:#1baf7a;--pos:#158a4a;--neg:#d03b3b;}
@media (prefers-color-scheme:dark){:root:where(:not([data-theme=light])) .viz{color-scheme:dark;
  --bg:#111110;--surface-1:#1a1a19;--border:#2c2c2a;--text-primary:#fff;
  --text-secondary:#c3c2b7;--muted:#8f8e86;--grid:#26261f;
  --income:#3987e5;--expense:#d95926;--cat:#3987e5;--cp:#199e70;--pos:#37c26f;--neg:#e66767;}}
:root[data-theme=dark] .viz{color-scheme:dark;
  --bg:#111110;--surface-1:#1a1a19;--border:#2c2c2a;--text-primary:#fff;
  --text-secondary:#c3c2b7;--muted:#8f8e86;--grid:#26261f;
  --income:#3987e5;--expense:#d95926;--cat:#3987e5;--cp:#199e70;--pos:#37c26f;--neg:#e66767;}
.viz{background:var(--bg);color:var(--text-primary);
  font:14px/1.5 -apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,sans-serif;min-height:100vh}
h1{font-size:20px;margin:0 0 2px} .sub{color:var(--muted);font-size:13px;margin:0 0 22px}
.tiles{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:12px;margin-bottom:22px}
.tile{background:var(--surface-1);border:1px solid var(--border);border-radius:12px;padding:14px 16px}
.tile .k{color:var(--muted);font-size:12px;text-transform:uppercase;letter-spacing:.04em}
.tile .v{font-size:22px;font-weight:650;margin-top:4px}
.card{background:var(--surface-1);border:1px solid var(--border);border-radius:14px;padding:18px 20px;margin-bottom:16px}
.card h2{font-size:14px;margin:0 0 14px;font-weight:600}
.card h2 .m{color:var(--muted);font-weight:400}
.grid2{display:grid;grid-template-columns:1fr 1fr;gap:16px}
@media(max-width:760px){.grid2{grid-template-columns:1fr}}
.chart,.hb{width:100%;height:auto;overflow:visible}
.grid{stroke:var(--grid);stroke-width:1}
.ytick{fill:var(--muted);font-size:10px;text-anchor:end}
.xtick{fill:var(--text-secondary);font-size:10px;text-anchor:middle}
.xnet{font-size:9px;text-anchor:middle}
.hbl{fill:var(--text-secondary);font-size:12px}
.hbv{fill:var(--text-primary);font-size:11px;font-weight:600}
.legend{display:flex;gap:16px;margin-top:8px;font-size:12px;color:var(--text-secondary)}
.legend span{display:inline-flex;align-items:center;gap:6px}
.legend i{width:10px;height:10px;border-radius:3px;display:inline-block}
table{width:100%;border-collapse:collapse;font-size:13px}
th,td{text-align:left;padding:7px 8px;border-bottom:1px solid var(--border)}
th{color:var(--muted);font-weight:500;font-size:11px;text-transform:uppercase;letter-spacing:.03em}
td.n,th.n{text-align:right;font-variant-numeric:tabular-nums}
.empty{color:var(--muted);font-style:italic}
footer{color:var(--muted);font-size:12px;margin-top:24px;line-height:1.7}
"""


def render_report_html(d, generated, nw=None):
    lt = d["lifetime"]
    net = lt["net"]
    sav = (net / lt["income"] * 100) if lt["income"] else 0
    dr = d["date_range"]
    tiles = ([("Net worth", inr_c(nw["net_worth"]))] if nw else []) + [
        ("Total balance", inr_c(d["total_balance"])),
        ("Lifetime in", inr_c(lt["income"])),
        ("Lifetime out", inr_c(lt["expense"])),
        ("Net", inr_c(net)),
        ("Savings rate", f"{sav:.0f}%"),
        ("Transactions", f'{d["n"]:,}'),
    ]
    tiles_html = "".join(f'<div class="tile"><div class="k">{_esc(k)}</div><div class="v">{_esc(v)}</div></div>' for k, v in tiles)

    rec_rows = "".join(
        f'<tr><td>{_esc(str(cp)[:26])}</td><td class="n">{mths}</td><td class="n">{cnt}</td><td class="n">{_esc(inr(amt))}</td></tr>'
        for cp, mths, amt, cnt in d["recurring"][:10]) or '<tr><td colspan="4" class="empty">None found.</td></tr>'
    big_rows = "".join(
        f'<tr><td>{_esc(date)}</td><td>{_esc(str(who)[:24])}</td><td>{_esc(cat)}</td><td class="n">{_esc(inr(amt))}</td></tr>'
        for date, who, cat, amt in d["biggest"][:10]) or '<tr><td colspan="4" class="empty">None.</td></tr>'

    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Fold Report</title><style>{_CSS}</style></head>
<body class="viz"><div class="wrap">
<h1>Fold · personal finance report</h1>
<p class="sub">Generated {_esc(generated)} · {_esc(dr[0])} → {_esc(dr[1])} · all data local to this machine</p>
<div class="tiles">{tiles_html}</div>
<div class="card"><h2>Monthly income vs expense <span class="m">· last 12 months (₹ net below each)</span></h2>
{_monthly_svg(d["monthly"])}
<div class="legend"><span><i style="background:var(--income)"></i>Income</span><span><i style="background:var(--expense)"></i>Expense</span></div></div>
{_investments_card(nw) if nw else ""}
<div class="grid2">
<div class="card"><h2>Spending by category <span class="m">· {_esc(d["month"])}</span></h2>{_hbars(d["cat_month"], "--cat")}</div>
<div class="card"><h2>Top counterparties <span class="m">· {_esc(d["month"])}</span></h2>{_hbars(d["merch_month"], "--cp")}</div>
</div>
<div class="grid2">
<div class="card"><h2>Recurring payments <span class="m">· seen ≥3 months</span></h2>
<table><thead><tr><th>Payee</th><th class="n">Months</th><th class="n">Txns</th><th class="n">Total</th></tr></thead><tbody>{rec_rows}</tbody></table></div>
<div class="card"><h2>Biggest expenses <span class="m">· all time</span></h2>
<table><thead><tr><th>Date</th><th>To</th><th>Category</th><th class="n">Amount</th></tr></thead><tbody>{big_rows}</tbody></table></div>
</div>
<footer>Most of your spend is P2P UPI to merchant VPAs with no merchant name, so category buckets are coarse — the <b>counterparty</b> and <b>recurring</b> views are the sharper lens. Balances update ~once/day (Fold's AA refresh).</footer>
</div></body></html>"""


# ---- spreadsheet export ----------------------------------------------------
def export_xlsx(db_path, out_path, aliases_path=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    aliases = load_aliases(aliases_path)
    d = compute(db_path, aliases_path=aliases_path)
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="20294C")

    def style_header(ws, ncols, money_cols=()):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in money_cols:
            for cell in ws[col][1:]:
                cell.number_format = '#,##0.00'

    wb = Workbook()
    tx = wb.active
    tx.title = "Transactions"
    tx.append(["Date", "Time(UTC)", "Type", "Mode", "Amount", "Cur",
               "Category (enriched)", "Counterparty", "Merchant", "Narration",
               "Balance", "Account", "Txn ID", "UUID"])
    db = sqlite3.connect(str(db_path))
    q = ("SELECT txn_date, txn_timestamp, type, mode, amount, currency, category_name, "
         "merchant_name, narration, current_balance, account_id, transaction_id, uuid "
         "FROM transactions ORDER BY txn_timestamp DESC")
    for (date, ts, typ, mode, amt, cur, cat, merch, narr, bal, acc, tid, u) in db.execute(q):
        cp, ecat = enrich(narr, merch, cat, typ, aliases)
        tx.append([date, (ts or "")[11:19], typ, mode, amt, cur, ecat, cp, merch,
                   narr, bal, acc, tid, u])
    db.close()
    style_header(tx, 14, money_cols=("E", "K"))
    for col, w in {"A": 11, "B": 10, "C": 10, "D": 8, "E": 12, "G": 20, "H": 22,
                   "I": 16, "J": 40, "K": 14}.items():
        tx.column_dimensions[col].width = w

    def sheet(name, headers, rows, money_cols=()):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for r in rows:
            ws.append(list(r))
        style_header(ws, len(headers), money_cols)
        return ws

    sheet("Monthly", ["Month", "Income", "Expense", "Net"],
          [[m["month"], m["income"], m["expense"], m["net"]] for m in d["monthly"]],
          money_cols=("B", "C", "D"))
    sheet("By Category", ["Category", "Total", "Txns"], d["cat_all"], money_cols=("B",))
    sheet("Top Payees", ["Counterparty", "Total", "Txns"], d["merch_all"], money_cols=("B",))
    sheet("Recurring", ["Payee", "Months", "Total", "Txns"],
          [[cp, mths, amt, cnt] for cp, mths, amt, cnt in d["recurring"]], money_cols=("C",))

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))
    return {"transactions": tx.max_row - 1, "sheets": wb.sheetnames}


# ---- extra analytics -------------------------------------------------------
def spend_by_mode(rows):
    from collections import defaultdict
    d = defaultdict(lambda: [0.0, 0])
    for r in rows:
        if r["type"] == "OUTGOING" and _spendable(r):
            d[r["mode"]][0] += r["amount"]; d[r["mode"]][1] += 1
    return sorted(([k, round(v[0], 2), v[1]] for k, v in d.items()), key=lambda x: -x[1])


def forecast_month(rows, month=None):
    import calendar
    month = month or dt.date.today().strftime("%Y-%m")
    y, m = (int(x) for x in month.split("-"))
    dim = calendar.monthrange(y, m)[1]
    today = dt.date.today()
    elapsed = today.day if (today.year, today.month) == (y, m) else dim
    spent = sum(r["amount"] for r in rows
                if r["month"] == month and r["type"] == "OUTGOING" and _spendable(r))
    prev = sorted({r["month"] for r in rows if r["month"] and r["month"] < month})[-3:]
    avg3 = (sum(r["amount"] for r in rows if r["month"] in prev and r["type"] == "OUTGOING"
                and _spendable(r)) / len(prev)) if prev else 0
    return {"month": month, "elapsed": elapsed, "days": dim, "spent": round(spent, 2),
            "projected": round(spent / elapsed * dim, 2) if elapsed else 0, "avg3": round(avg3, 2)}


def tax_year(db_path, fy=None, aliases_path=None):
    from collections import defaultdict
    rows, _ = _load(db_path, aliases_path)
    if fy:
        y1 = int(fy.split("-")[0])
    else:
        t = dt.date.today(); y1 = t.year if t.month >= 4 else t.year - 1
    start, end = f"{y1}-04-01", f"{y1 + 1}-03-31"
    sel = [r for r in rows if start <= r["date"] <= end and _spendable(r)]
    inc = sum(r["amount"] for r in sel if r["type"] == "INCOMING")
    exp = sum(r["amount"] for r in sel if r["type"] == "OUTGOING")
    cat = defaultdict(float)
    mon = defaultdict(lambda: [0.0, 0.0])
    for r in sel:
        if r["type"] == "OUTGOING":
            cat[r["category"]] += r["amount"]; mon[r["month"]][1] += r["amount"]
        else:
            mon[r["month"]][0] += r["amount"]
    return {"fy": f"{y1}-{str(y1 + 1)[2:]}", "start": start, "end": end, "n": len(sel),
            "income": round(inc, 2), "expense": round(exp, 2), "net": round(inc - exp, 2),
            "by_category": sorted(([k, round(v, 2)] for k, v in cat.items()), key=lambda x: -x[1]),
            "monthly": [[m, round(v[0], 2), round(v[1], 2)] for m, v in sorted(mon.items())]}


def compare_months(db_path, m1, m2, aliases_path=None):
    from collections import defaultdict
    rows, _ = _load(db_path, aliases_path)

    def agg(m):
        inc = sum(r["amount"] for r in rows if r["month"] == m and r["type"] == "INCOMING" and _spendable(r))
        exp = sum(r["amount"] for r in rows if r["month"] == m and r["type"] == "OUTGOING" and _spendable(r))
        c = defaultdict(float)
        for r in rows:
            if r["month"] == m and r["type"] == "OUTGOING" and _spendable(r):
                c[r["category"]] += r["amount"]
        return inc, exp, c
    i1, e1, c1 = agg(m1); i2, e2, c2 = agg(m2)
    cats = sorted(set(c1) | set(c2), key=lambda k: -(c2.get(k, 0) + c1.get(k, 0)))
    return {"m1": m1, "m2": m2, "income": [round(i1, 2), round(i2, 2)],
            "expense": [round(e1, 2), round(e2, 2)],
            "by_category": [[k, round(c1.get(k, 0), 2), round(c2.get(k, 0), 2),
                             round(c2.get(k, 0) - c1.get(k, 0), 2)] for k in cats]}


def unusual(db_path, mult=2.5, floor=1000, aliases_path=None):
    from collections import defaultdict
    from statistics import mean
    rows, _ = _load(db_path, aliases_path)
    groups = defaultdict(list)
    for r in rows:
        if r["type"] == "OUTGOING" and _spendable(r) and r["counterparty"]:
            groups[r["counterparty"]].append(r)
    flagged = []
    for cp, rs in groups.items():
        if len(rs) < 3:
            continue
        avg = mean(x["amount"] for x in rs)
        for x in rs:
            if x["amount"] >= max(floor, mult * avg):
                flagged.append([x["date"], cp, round(x["amount"], 2), round(avg, 2)])
    return sorted(flagged, key=lambda z: -z[2])[:20]


def alias_suggest(db_path, path, top=40):
    rows, _ = _load(db_path, None)
    from collections import defaultdict
    d = defaultdict(float)
    for r in rows:
        if r["type"] == "OUTGOING" and _spendable(r) and r["category"] == "Uncategorized" and r["counterparty"]:
            d[r["counterparty"]] += r["amount"]
    pp = Path(path)
    existing = json.loads(pp.read_text()) if pp.exists() else {}
    existing.setdefault("_help", "pattern (matched in counterparty/narration, lowercase) -> {name, category}")
    added = 0
    for cp, _amt in sorted(d.items(), key=lambda x: -x[1])[:top]:
        if cp.lower() not in existing:
            existing[cp.lower()] = {"name": "", "category": ""}
            added += 1
    pp.write_text(json.dumps(existing, indent=2, ensure_ascii=False))
    return added


def export_lunchmoney_csv(db_path, out_path, aliases_path=None):
    import csv as _csv
    aliases = load_aliases(aliases_path)
    db = sqlite3.connect(str(db_path))
    q = ("SELECT txn_date, type, amount, currency, category_name, merchant_name, narration "
         "FROM transactions ORDER BY txn_timestamp DESC")
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    n = 0
    with open(out_path, "w", newline="") as f:
        w = _csv.writer(f)
        w.writerow(["date", "payee", "amount", "currency", "category", "notes"])
        for date, typ, amt, cur, cat, merch, narr in db.execute(q):
            cp, ecat = enrich(narr, merch, cat, typ, aliases)
            signed = (amt or 0) if typ == "INCOMING" else -(amt or 0)
            w.writerow([date, cp or (narr or "")[:40], f"{signed:.2f}", cur or "INR", ecat, narr])
            n += 1
    db.close()
    return n
